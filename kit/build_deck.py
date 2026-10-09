#!/usr/bin/env python3
"""Fill deck-template.pptx from demo-config.json -> deck.pptx  (needs: pip install python-pptx)"""
import json, re, pathlib
from pptx import Presentation
d = pathlib.Path(__file__).parent
cfg = json.loads((d/'demo-config.json').read_text())
S, ser, ini = cfg['story'], cfg['series'], cfg['tracker']['initiatives']
b = cfg['brand']
def fmt(x): return ('%g' % round(x, 1))
vb, va = ser['velocity']['before'], ser['velocity']['after']
sc = S['a5']['scorecard']
quote = S['a5']['quote']
body, _, attr = quote.partition(' — ')
sents = re.split(r'(?<=[.?!])\s+(?=[A-Z])', body.strip())
RAG = {'g': 'Green', 'y': 'Yellow', 'r': 'Red'}
cfg['calc'] = {
  'title_line': f"{b['title_a']} {b['title_b']}",
  'definition': f"{b['definition_term']} (adj.) — {b['definition']}",
  'a1_quote_short': '“' + S['a1']['quote'].strip('"').split('?')[0] + '?”',
  'a2_quote': '“' + S['a2']['quote'].strip('"') + '”',
  'tracker_sub': 'Manually maintained · ' + cfg['tracker']['stale_tag'].lower(),
  'before_line': 'Before:  ' + ' + '.join(map(fmt, vb)) + f'  =  {fmt(sum(vb))} days',
  'after_line':  'After:  ' + ' + '.join(map(fmt, va)) + f'  =  {fmt(sum(va))} days',
  'quote1': '“' + sents[0].strip('"“”'), 'quote2': ' '.join(sents[1:]).strip('"“”') + '”',
  'quote_attr': '— ' + attr.strip().rstrip('.'),
}
for i, r in enumerate(ini):
    cfg['calc'][f'rag{i}'] = RAG[r['rag']]
    cfg['calc'][f'note{i}'] = r['note'].replace('"', '“', 1).replace('"', '”', 1) if r.get('strike') else r['note']
for i, st in enumerate(S['a4']['stats']): cfg['calc'][f'a4_num{i}'] = st['num'].replace('&times;', '×')
for i, r in enumerate(S['a5']['deck_cards']):
    cfg['calc'][f'a5_val{i}'] = r['value']; cfg['calc'][f'a5_note{i}'] = r['label']
def get(path):
    cur = cfg
    for k in path.split('.'):
        cur = cur[int(k)] if isinstance(cur, list) else cur[k]
    return str(cur).replace('&plusmn;', '±').replace('&amp;', '&')
def fill(par):
    if not par.runs or '{{' not in par.text: return
    txt = re.sub(r'\{\{([^}]+)\}\}', lambda m: get(m.group(1)), par.text)
    par.runs[0].text = txt
    for r in par.runs[1:]: r.text = ''
p = Presentation(d/'deck-template.pptx')
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for par in sh.text_frame.paragraphs: fill(par)
# footer note + attribution on every slide's notes (keeps slide visuals clean)
for s in p.slides:
    s.notes_slide.notes_text_frame.text = cfg['footer']['disclaimer'] + ' ' + cfg['footer']['attribution']
p.save(d/'deck.pptx'); print('wrote', d/'deck.pptx')
