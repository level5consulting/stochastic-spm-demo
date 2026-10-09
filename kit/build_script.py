#!/usr/bin/env python3
"""Fill script-template.docx from demo-config.json -> run-of-show.docx (needs: pip install python-docx).
The structural fields (title, act timings) come from the config; the spoken narration is authored prose in the template - edit it in Word."""
import json, re, pathlib
from docx import Document
d = pathlib.Path(__file__).parent
cfg = json.loads((d/'demo-config.json').read_text())
b = cfg['brand']
cfg['calc'] = {'brand_upper': (b['title_a'] + ' ' + b['title_b']).upper()}
def get(path):
    cur = cfg
    for k in path.split('.'): cur = cur[int(k)] if isinstance(cur, list) else cur[k]
    return str(cur)
def sub(p):
    for r in p.runs:
        if '{{' in r.text: r.text = re.sub(r'\{\{([^}]+)\}\}', lambda m: get(m.group(1)), r.text)
doc = Document(d/'script-template.docx')
for p in doc.paragraphs: sub(p)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: sub(p)
doc.save(d/'run-of-show.docx'); print('wrote', d/'run-of-show.docx')
